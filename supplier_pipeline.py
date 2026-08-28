#!/usr/bin/env python3
"""
supplier_pipeline.py
====================

Take a supplier xlsx/csv, generate spec JSONs via Gemini, then link each spec
to the matching unmatched SKU on the given supplier.

Usage
-----
    export GEMINI_API_KEY=...
    python supplier_pipeline.py path/to/file.xlsx --supplier-id <UUID>

Matching key is the supplier's own SKU code (== Cerve `supplier_sku_id`).
Auto-detects the column from common names; pass --sku-column to override.

In-process Gemini calls (no subprocess) so the whole thing can be frozen
with PyInstaller.
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import sys
from pathlib import Path
from typing import Any

import jsonschema
import requests

import gemini_csv_to_specs
import spec_cleanup

HERE = Path(__file__).resolve().parent


def _bundle_dir() -> Path:
    """Where bundled data files live (ExtractSpec.txt, spec_schema.json).
    sys._MEIPASS is PyInstaller's extracted-archive directory; in dev we use
    the script's own folder."""
    return Path(getattr(sys, "_MEIPASS", HERE))


def _user_data_dir() -> Path:
    """Where the app writes its outputs (specs_out/, normalised.csv).
    In a PyInstaller bundle we don't write next to the binary — that folder
    is read-only (and on --onefile, ephemeral)."""
    if getattr(sys, "frozen", False):
        return Path.home() / "Documents" / "SupplierPipeline"
    return HERE


PROMPT_FILE = _bundle_dir() / "ExtractSpec.txt"
SCHEMA_FILE = _bundle_dir() / "spec_schema.json"

AUTH_URL = "https://auth.cerve.com/v2/token"
SPEC_URL = "https://suppliers.cerve.com/v2/specs"
SUPPLIERS_BASE = "https://suppliers.cerve.com/v2/suppliers"

SKU_COL_CANDIDATES = [
    "supplier_sku_id", "SKU ID*", "Supplier Code", "Product Code",
    "RY Code", "SKU code", "sku_id",
]


# ---------------------------------------------------------------------------
# File reading (csv / xlsx / xls -> normalised csv)
# ---------------------------------------------------------------------------
def stringify(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# Words that mark a row as column headers rather than data.
HEADER_KEYWORDS = {
    "sku", "code", "gtin", "ean", "barcode", "product", "description", "name",
    "brand", "pack", "case", "price", "rrp", "rsp", "weight", "volume", "qty",
    "quantity", "vat", "supplier", "ry",
}

HEADER_SCAN_DEPTH = 30


def _score_header_row(cells: list[str]) -> tuple[int, int]:
    """(keyword count, populated count). Higher is better. Tie-break by row order."""
    populated = sum(1 for c in cells if c)
    kw = 0
    for c in cells:
        cl = c.lower()
        if any(k in cl for k in HEADER_KEYWORDS):
            kw += 1
    return kw, populated


def _pick_header_row(rows_iter: list[list[str]]) -> int | None:
    """Pick the row index whose cells look most like column headers.
    Prefer the row with most header keywords; fall back to first row with >= 4
    populated cells if no row has any keyword matches. Returns None if neither."""
    best_idx = None
    best_score = (0, 0)
    for i, cells in enumerate(rows_iter[:HEADER_SCAN_DEPTH]):
        score = _score_header_row(cells)
        if score[0] >= 1 and score > best_score:
            best_score = score
            best_idx = i
    if best_idx is not None:
        return best_idx
    # No keyword hits anywhere — fall back to the first reasonably wide row.
    for i, cells in enumerate(rows_iter[:HEADER_SCAN_DEPTH]):
        if sum(1 for c in cells if c) >= 4:
            return i
    return None


def _prune_empty_columns(headers: list[str], rows: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    """Drop columns where the header is blank — these are trailing pads added
    by openpyxl when a sheet's max_column extends past the real data."""
    keep_idx = [i for i, h in enumerate(headers) if h]
    new_headers = [headers[i] for i in keep_idx]
    new_rows = [[row[i] if i < len(row) else "" for i in keep_idx] for row in rows]
    return new_headers, new_rows


def read_file(src: Path) -> tuple[list[str], list[dict[str, str]]]:
    """Read csv/xlsx/xls and return (headers, rows). No filtering, no writes."""
    ext = src.suffix.lower()
    if ext == ".csv":
        with src.open(newline="", encoding="utf-8-sig") as fh:
            reader = csv.DictReader(fh)
            headers = [(h or "").strip() for h in (reader.fieldnames or [])]
            row_cells = [[stringify(r.get(h, "")) for h in headers] for r in reader]
    elif ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        all_rows = [
            [stringify(v) for v in r] for r in ws.iter_rows(values_only=True)
        ]
        idx = _pick_header_row(all_rows)
        if idx is None:
            raise SystemExit("no header row found in spreadsheet")
        headers = all_rows[idx]
        row_cells = all_rows[idx + 1:]
    elif ext == ".xls":
        import xlrd
        wb = xlrd.open_workbook(str(src))
        ws = wb.sheets()[0]
        all_rows = [[stringify(v) for v in ws.row_values(i)] for i in range(ws.nrows)]
        idx = _pick_header_row(all_rows)
        if idx is None:
            raise SystemExit("no header row found in spreadsheet")
        headers = all_rows[idx]
        row_cells = all_rows[idx + 1:]
    else:
        raise SystemExit(f"unsupported file type: {ext}")

    headers, row_cells = _prune_empty_columns(headers, row_cells)
    rows = [{headers[i]: row[i] for i in range(len(headers)) if i < len(row)} for row in row_cells]
    return headers, rows


def peek_headers(src: Path) -> list[str]:
    """Just the headers — used by the UI to populate the SKU-column dropdown."""
    return read_file(src)[0]


def read_to_csv(src: Path, dst: Path) -> tuple[list[str], list[dict[str, str]]]:
    headers, rows = read_file(src)

    # Drop rows that carry almost nothing — blank lines and section dividers.
    def keep(row: dict[str, str]) -> bool:
        return sum(1 for v in row.values() if v) >= 2

    kept = [r for r in rows if keep(r)]

    with dst.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=headers, extrasaction="ignore")
        w.writeheader()
        w.writerows(kept)

    print(f"[READ] {len(rows)} rows -> {len(kept)} kept -> {dst}")
    return headers, kept


def auto_detect_sku_column(headers: list[str], extra_candidates: list[str] | None = None) -> str | None:
    """Return the first header that matches a known SKU-column candidate, else None.
    `extra_candidates` is a caller-supplied list (e.g. user-added names persisted by the UI)."""
    candidates = SKU_COL_CANDIDATES + list(extra_candidates or [])
    lower = {h.lower(): h for h in headers}
    for c in candidates:
        if c.lower() in lower:
            return lower[c.lower()]
    return None


def find_sku_column(headers: list[str], explicit: str | None) -> str:
    if explicit:
        if explicit not in headers:
            raise SystemExit(f"column '{explicit}' not in headers: {headers}")
        return explicit
    lower = {h.lower(): h for h in headers}
    for c in SKU_COL_CANDIDATES:
        if c.lower() in lower:
            return lower[c.lower()]
    raise SystemExit(f"could not auto-detect SKU code column; pass --sku-column. headers: {headers}")


# ---------------------------------------------------------------------------
# Spec generation (in-process Gemini)
# ---------------------------------------------------------------------------
def generate_specs(csv_path: Path, out_dir: Path) -> None:
    if not os.environ.get("GEMINI_API_KEY") and not os.environ.get("GOOGLE_API_KEY"):
        raise SystemExit("set GEMINI_API_KEY (https://aistudio.google.com/apikey)")
    gemini_csv_to_specs.process_csv(
        csv_path=csv_path,
        prompt_file=PROMPT_FILE,
        schema_file=SCHEMA_FILE,
        output_dir=out_dir,
    )


# ---------------------------------------------------------------------------
# Cerve client
# ---------------------------------------------------------------------------
class CerveClient:
    def __init__(self, client_id: str, client_secret: str):
        self.client_id, self.client_secret = client_id, client_secret
        self.token = self._auth()

    def _auth(self) -> str:
        r = requests.post(
            AUTH_URL,
            headers={"Content-Type": "application/x-www-form-urlencoded"},
            data={"grant_type": "client_credentials",
                  "client_id": self.client_id, "client_secret": self.client_secret},
            timeout=60,
        )
        r.raise_for_status()
        print("[AUTH] ok")
        return r.json()["access_token"]

    def request(self, method: str, url: str, **kw) -> requests.Response:
        headers = dict(kw.pop("headers", {}) or {})
        headers["Authorization"] = f"Bearer {self.token}"
        kw.setdefault("timeout", 60)
        r = requests.request(method, url, headers=headers, **kw)
        if r.status_code == 401:
            print("[AUTH] 401 — refreshing")
            self.token = self._auth()
            headers["Authorization"] = f"Bearer {self.token}"
            r = requests.request(method, url, headers=headers, **kw)
        return r


def get_credentials() -> tuple[str, str]:
    cid = os.environ.get("CERVE_CLIENT_ID")
    sec = os.environ.get("CERVE_CLIENT_SECRET")
    if cid and sec:
        return cid, sec
    sys.path.insert(0, str(HERE))
    from Authentication import AuthenticationRequirements  # type: ignore
    return AuthenticationRequirements.CLIENT_ID, AuthenticationRequirements.SECRET


# ---------------------------------------------------------------------------
# Schema validation
# ---------------------------------------------------------------------------
_SCHEMA_CACHE: dict | None = None


def _strip_required(obj):
    """Recursively remove `required` lists from the schema. The Cerve schema
    marks fields as required that are routinely null at create-time (and which
    we strip from the outbound payload), so leaving `required` in causes false
    positives. We keep type/enum/oneOf/etc. — the actual shape constraints —
    so the validator still catches the real bugs (wrong types, bad enums,
    malformed structures). Cerve enforces required-field presence on POST."""
    if isinstance(obj, dict):
        obj.pop("required", None)
        for v in obj.values():
            _strip_required(v)
    elif isinstance(obj, list):
        for v in obj:
            _strip_required(v)


def _load_schema() -> dict:
    """Read spec_schema.json, recursively strip `required`, cache the result."""
    global _SCHEMA_CACHE
    if _SCHEMA_CACHE is None:
        s = json.loads(SCHEMA_FILE.read_text(encoding="utf-8"))
        _strip_required(s)
        _SCHEMA_CACHE = s
    return _SCHEMA_CACHE


def _strip_nulls_deep(obj):
    """Recursive null strip used ONLY by validate_spec. We send specs to Cerve
    with nested nulls intact (Cerve requires structural keys like product.weight
    even when null), but the schema marks those fields as e.g. "string" with no
    null variant — so we strip nulls deeply *for validation* to avoid false
    positives on server-managed fields like timestamps.cerve.updated_at."""
    if isinstance(obj, dict):
        return {k: _strip_nulls_deep(v) for k, v in obj.items() if v is not None}
    if isinstance(obj, list):
        return [_strip_nulls_deep(v) for v in obj if v is not None]
    return obj


def validate_spec(spec: dict) -> str | None:
    """Check a spec dict before POSTing. Returns None if valid, else a one-line
    error message naming the failing field/path."""
    if not spec.get("name"):
        return "missing required field: name"
    if not spec.get("unit_of_measure"):
        return "missing required field: unit_of_measure"

    try:
        jsonschema.validate(instance=_strip_nulls_deep(spec), schema=_load_schema())
    except jsonschema.exceptions.ValidationError as e:
        path = ".".join(str(p) for p in e.absolute_path) or "/"
        return f"schema error at {path}: {e.message}"
    return None


# ---------------------------------------------------------------------------
# Cerve actions
# ---------------------------------------------------------------------------
def list_unmatched_skus(client: CerveClient, supplier_id: str) -> dict[str, str]:
    """Return {supplier_sku_id: cerve_sku_id} for SKUs with no spec linked."""
    url = f"{SUPPLIERS_BASE}/{supplier_id}/skus"
    out: dict[str, str] = {}
    token = None
    while True:
        params = {"page_size": "1000", "spec_id__is_null": "true"}
        if token:
            params["page_token"] = token
        r = client.request("GET", url, params=params)
        if r.status_code != 200:
            raise SystemExit(f"list SKUs failed: {r.status_code} {r.text[:300]}")
        body = r.json()
        for sku in body.get("skus") or []:
            sid, cid = sku.get("supplier_sku_id"), sku.get("id")
            if sid and cid:
                out[sid] = cid
        token = body.get("next_page_token")
        if not token:
            break
    print(f"[SKU LIST] {len(out)} unmatched SKUs on supplier")
    return out


def _strip_nulls(obj):
    """Drop only TOP-LEVEL null fields. Cerve's POST /v2/specs rejects null
    top-level fields like `content_declarations: null` → 400, but it REQUIRES
    nested structural keys (e.g. product.weight) to be present even when their
    values are null. So we strip nulls only at the spec root and let nested
    null/empty structure survive — spec_cleanup is responsible for ensuring the
    nested keys exist with the right shape."""
    if isinstance(obj, dict):
        return {k: v for k, v in obj.items() if v is not None}
    return obj


def post_spec(client: CerveClient, spec: dict) -> tuple[str | None, str, dict | None]:
    """Create the spec on Cerve.

    Returns ``(spec_id, status, failure)`` where ``status`` is:
      - ``"created"`` — POST succeeded, new spec; failure is None.
      - ``"reused"``  — POST returned 409 (duplicate GTIN); the returned id
                         points at an EXISTING spec that we did NOT modify;
                         failure is None.
      - ``"failed"``  — spec_id is None; ``failure`` is a dict with
                         ``{"status": int, "response": str, "payload": dict,
                            "reason": str}`` that the caller can log + save.
    """
    spec = _strip_nulls(spec)
    r = client.request("POST", SPEC_URL,
                       headers={"Content-Type": "application/json"}, json=spec)
    if r.status_code in (200, 201):
        return r.json().get("id"), "created", None
    if r.status_code == 409:
        # Duplicate by GTIN — find the existing one so we can still link.
        meta = (r.json().get("details") or [{}])[0].get("metadata") or {}
        gtin = meta.get("ean_13") or meta.get("upc_a") or meta.get("ean_8")
        if gtin:
            look = client.request("GET", SPEC_URL, params={"page_size": 1, "query": gtin})
            if look.status_code == 200:
                specs = look.json().get("specs") or []
                if specs:
                    return specs[0]["id"], "reused", None
        return None, "failed", {
            "status": 409,
            "response": r.text,
            "payload": spec,
            "reason": f"duplicate GTIN ({gtin}) but lookup returned no spec",
        }

    reason = f"HTTP {r.status_code}"
    try:
        details = (r.json().get("details") or [{}])[0]
        md = details.get("metadata") or {}
        if md:
            first_field, first_msg = next(iter(md.items()))
            reason = f"{first_field}: {first_msg}"
        elif details.get("reason"):
            reason = details["reason"]
    except (ValueError, KeyError, StopIteration):
        pass
    return None, "failed", {
        "status": r.status_code,
        "response": r.text,
        "payload": spec,
        "reason": reason,
    }


def _process_inline_sub_specs(parent_spec: dict, client: CerveClient,
                              dropped_gtins_log: list, supplier_sku_id: str) -> tuple[bool, str | None]:
    """When the model emits a two-level pack (e.g. ``10x5x20g``), it nests the
    child spec inline at ``parent.sub_specs[i]._inline_spec``. The child must
    exist on Cerve before the parent can be POSTed with a real ``spec_id``, so
    we clean+validate+POST each inline child here and replace the marker with
    the returned id. Returns ``(ok, err)``."""
    sub_specs = parent_spec.get("sub_specs")
    if not sub_specs:
        return True, None
    for entry in sub_specs:
        if not isinstance(entry, dict):
            continue
        inline = entry.pop("_inline_spec", None)
        if inline is None:
            continue

        try:
            child_cleaned, child_dropped = spec_cleanup.clean_spec(inline)
            for d in child_dropped:
                dropped_gtins_log.append({
                    "supplier_sku_id": f"{supplier_sku_id} (child)",
                    "spec_name": child_cleaned.get("name", ""),
                    "field": d["field"],
                    "value": d["value"],
                })
        except Exception as e:  # noqa: BLE001
            return False, f"child cleanup failed: {type(e).__name__}: {e}"

        err = validate_spec(child_cleaned)
        if err:
            return False, f"child invalid — {err}"

        child_id, child_status, child_fail = post_spec(client, child_cleaned)
        if child_id is None:
            reason = (child_fail or {}).get("reason", "?")
            return False, f"child POST failed — {reason}"
        verb = "created" if child_status == "created" else "reused existing"
        print(f"  [sub-spec] {verb} child spec {child_id} "
              f"(name={child_cleaned.get('name')!r}, quantity={entry.get('quantity')})")
        entry["spec_id"] = child_id
    return True, None


def link_sku(client: CerveClient, supplier_id: str, sku_id: str, spec_id: str) -> bool:
    r = client.request(
        "PATCH", f"{SUPPLIERS_BASE}/{supplier_id}/skus/{sku_id}",
        headers={"Content-Type": "application/json"},
        json={"spec_id": spec_id},
    )
    if r.status_code != 200:
        print(f"[LINK] {sku_id} failed {r.status_code}: {r.text[:200]}")
        return False
    return True


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    ap.add_argument("supplier_file", type=Path, help="xlsx, xls, or csv")
    ap.add_argument("--supplier-id", required=True, help="Cerve supplier UUID")
    ap.add_argument("--sku-column", help="Override auto-detection of supplier SKU code column")
    args = ap.parse_args()

    if not args.supplier_file.is_file():
        raise SystemExit(f"file not found: {args.supplier_file}")

    out_dir = _user_data_dir() / "specs_out" / args.supplier_file.stem
    out_dir.mkdir(parents=True, exist_ok=True)
    normalised_csv = out_dir / "normalised.csv"

    # 1. Read + normalise
    headers, rows = read_to_csv(args.supplier_file, normalised_csv)
    sku_col = find_sku_column(headers, args.sku_column)
    print(f"[COLS] using sku column '{sku_col}'")

    # 2. Auth + list unmatched SKUs on the supplier
    client = CerveClient(*get_credentials())
    sku_lookup = list_unmatched_skus(client, args.supplier_id)

    # 3. Keep only rows whose SKU is actually waiting for a spec.
    filtered_rows = [r for r in rows if (r.get(sku_col) or "").strip() in sku_lookup]
    print(f"[FILTER] {len(rows)} rows -> {len(filtered_rows)} match an unmatched SKU")
    if not filtered_rows:
        print("[DONE] nothing to do — every row's SKU already has a spec")
        return 0

    # 4. Rewrite the normalised csv with just those rows, then generate specs.
    with normalised_csv.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=headers, extrasaction="ignore")
        w.writeheader()
        w.writerows(filtered_rows)

    generate_specs(normalised_csv, out_dir)
    for idx, row in enumerate(filtered_rows, start=1):
        sku_id = (row.get(sku_col) or "").strip()
        if not sku_id:
            continue
        for suffix in (".json", ".error.json"):
            src = out_dir / f"row_{idx}{suffix}"
            if src.is_file():
                src.rename(out_dir / f"{sku_id}{suffix}")

    # 5. For each row, clean + validate + POST the spec, then link the SKU.
    summary = {
        "linked_to_new_spec": 0,
        "linked_to_existing_spec": 0,
        "spec_only": 0,
        "no_spec": 0,
        "spec_failed": 0,
    }
    dropped_gtins_log: list[dict[str, str]] = []
    for idx, row in enumerate(filtered_rows, start=1):
        supplier_sku_id = (row.get(sku_col) or "").strip()
        spec_path = out_dir / f"{supplier_sku_id}.json"
        if not spec_path.is_file():
            print(f"[ROW {idx}] no spec json for sku '{supplier_sku_id}' — skip")
            summary["no_spec"] += 1
            continue
        spec = json.loads(spec_path.read_text(encoding="utf-8"))
        if not spec.get("name"):
            print(f"[ROW {idx}] empty spec (divider?) — skip")
            summary["no_spec"] += 1
            continue

        # Inline sub-specs must exist on Cerve before the parent is POSTed.
        ok, sub_err = _process_inline_sub_specs(spec, client, dropped_gtins_log, supplier_sku_id)
        if not ok:
            print(f"[ROW {idx}] sub-spec creation failed — {sub_err}")
            summary["spec_failed"] += 1
            continue

        # Local cleanup: GTIN validation, unit normalisation, allergen template.
        try:
            spec, dropped_gtins = spec_cleanup.clean_spec(spec)
            for d in dropped_gtins:
                dropped_gtins_log.append({
                    "supplier_sku_id": supplier_sku_id,
                    "spec_name": spec.get("name", ""),
                    "field": d["field"],
                    "value": d["value"],
                })
        except Exception as e:  # noqa: BLE001
            print(f"[ROW {idx}] spec cleanup failed, sending raw: {type(e).__name__}: {e}")

        validation_err = validate_spec(spec)
        if validation_err:
            print(f"[ROW {idx}] spec invalid — {validation_err}")
            summary["spec_failed"] += 1
            continue

        spec_id, spec_status, fail = post_spec(client, spec)
        if spec_status == "created":
            print(f"[ROW {idx}] created new spec {spec_id}")
        elif spec_status == "reused":
            print(f"[ROW {idx}] reused existing spec {spec_id} (matched by GTIN — not modified)")
        else:
            assert fail is not None
            spec_name = spec.get("name") or "<no name>"
            fail_path = out_dir / f"{supplier_sku_id}.post_failed.json"
            fail_path.write_text(
                json.dumps(fail, indent=2, ensure_ascii=False) + "\n",
                encoding="utf-8",
            )
            print(
                f"[ROW {idx}] POST {fail['status']} for '{spec_name}' (sku "
                f"{supplier_sku_id}) — {fail['reason']}"
            )
            print(f"[ROW {idx}] full payload + response: {fail_path.name}")
            summary["spec_failed"] += 1
            continue

        cerve_sku_id = sku_lookup[supplier_sku_id]
        if link_sku(client, args.supplier_id, cerve_sku_id, spec_id):
            print(f"[ROW {idx}] linked SKU {cerve_sku_id} -> spec {spec_id}")
            if spec_status == "created":
                summary["linked_to_new_spec"] += 1
            else:
                summary["linked_to_existing_spec"] += 1
        else:
            summary["spec_only"] += 1

    print(f"[DONE] {summary}")

    # 6. Report any GTINs the cleaner had to drop.
    if dropped_gtins_log:
        gtin_csv = out_dir / "dropped_gtins.csv"
        with gtin_csv.open("w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=["supplier_sku_id", "spec_name", "field", "value"])
            w.writeheader()
            w.writerows(dropped_gtins_log)
        unique_skus = len({d["supplier_sku_id"] for d in dropped_gtins_log})
        print(f"[GTIN] dropped {len(dropped_gtins_log)} invalid GTIN(s) across "
              f"{unique_skus} SKU(s) — see {gtin_csv.name}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
